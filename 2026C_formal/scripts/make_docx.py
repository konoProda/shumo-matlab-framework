#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_docx.py —— 把纯文本说明转成交付用的 .docx

用法：python3 make_docx.py <输入.txt> <输出.docx> <标题>

输入文本约定：以 '#' 开头的行为小标题，其余非空行为正文段落。
按交付规范，正文为**自然段落**，只写清三件事：算的是什么 / 按什么思路实现 / 得到什么结果；
不出现运行入口、命令行、目录树。
"""
import sys
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn


def set_font(doc):
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def main():
    src, dst, title = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(src, encoding='utf-8') as f:
        lines = [ln.rstrip() for ln in f]

    doc = Document()
    set_font(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.6)

    h = doc.add_heading(title, level=1)
    h.runs[0].font.name = 'Times New Roman'
    h.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    for ln in lines:
        t = ln.strip()
        if not t:
            continue
        if t.startswith('#'):
            p = doc.add_heading(t.lstrip('#').strip(), level=2)
            p.runs[0].font.name = 'Times New Roman'
            p.runs[0]._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        else:
            doc.add_paragraph(t)

    doc.save(dst)
    print(f'已生成 {dst}')


if __name__ == '__main__':
    main()
