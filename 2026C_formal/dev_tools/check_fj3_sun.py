# -*- coding: utf-8 -*-
"""带阈值的日出/日落时刻配对（阈值化以剔除夜间微小非零值）"""
import openpyxl
from collections import Counter

wb2 = openpyxl.load_workbook('data/附件/附件2.xlsx', data_only=True)
ws2 = wb2['光伏发电实际功率']
D, T = 365, 144
raw = [[float(ws2.cell(r, c).value) for c in range(2, 2+T)] for r in range(2, 2+D)]
pv = [[raw[0][T-1]] + raw[0][:T-1]] + [[raw[d-1][T-1]] + raw[d-1][:T-1] for d in range(1, D)]
act = [[sum(pv[d][6*h:6*h+6])/6.0 for h in range(24)] for d in range(D)]
nz = sorted(v for d in range(D) for h in (0,1,2,3,23) for v in [act[d][h]] if v > 0)
print('夜间(0-3时/23时) 实际光伏 >0 的取值：n=%d  min=%.4f  中位=%.4f  max=%.4f' % (len(nz), nz[0], nz[len(nz)//2], nz[-1]))

wb3 = openpyxl.load_workbook('data/附件/附件3.xlsx', data_only=True)
ws3 = wb3['Sheet1']
iss = [(ws3.cell(r, 2).value, [float(ws3.cell(r, c).value or 0) for c in range(3, 27)])
       for r in range(2, ws3.max_row + 1)]

THR = 10.0
print('\n阈值 = %.0f kW' % THR)
for lab, t0 in [('0:00', 0), ('6:00', 6), ('12:00', 12), ('18:00', 18)]:
    cA = cB = 0
    for d in range(D):
        al = [h for h in range(24) if act[d][h] > THR]
        if not al: continue
        fsun, fset = al[0], al[-1]
        fy = iss[d*4 + ['0:00','6:00','12:00','18:00'].index(lab)][1]
        k_ = [k for k in range(1, 25) if fy[k-1] > THR]
        if not k_: continue
        # 读A下预报的首个>阈小时 = t0+k-1；读B下 = t0+k
        da = abs((t0 + k_[0] - 1) - fsun); db = abs((t0 + k_[0]) - fsun)
        if da < db: cA += 1
        elif db < da: cB += 1
    print('  发布 %-5s  日出：读A更近 %3d 天 / 读B更近 %3d 天' % (lab, cA, cB))

print('\n【全年平均日曲线：只取"实际小时均值>10"的样本】')
print('%4s %10s %12s %12s' % ('小时', '实际', '预报(读A)', '预报(读B)'))
for h in range(24):
    A_ = [act[d][h] for d in range(D) if act[d][h] > THR]
    FA = [iss[d*4][1][h]     for d in range(D-1) if act[d][h] > THR]
    FB = [iss[d*4][1][h-1] for d in range(1, D) if act[d][h] > THR]
    print('%4d %10.1f %12.1f %12.1f   (n=%3d/%3d)' % (h, sum(A_)/max(len(A_),1),
          sum(FA)/max(len(FA),1), sum(FB)/max(len(FB),1), len(FA), len(FB)))
