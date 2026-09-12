# -*- coding: utf-8 -*-
"""附件3 结构与完整性校验（含数值列非数值文本扫描）"""
import openpyxl, datetime, collections

wb = openpyxl.load_workbook('data/附件/附件3.xlsx', data_only=True)
ws = wb['Sheet1']
print('dims', ws.dimensions, ws.max_row, ws.max_column)

hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print('表头前2:', hdr[:2], ' 末2:', hdr[-2:])
assert hdr[0] == '日期' and hdr[1] == '预报时刻'
assert hdr[2:] == ['预报%d小时' % k for k in range(1, 25)], hdr[2:]

times = ['0:00', '6:00', '12:00', '18:00']
rows = []
nondate_label = 0
text_cells = collections.Counter()
hours_by_slot = collections.defaultdict(list)

for r in range(2, ws.max_row + 1):
    a = ws.cell(r, 1).value
    b = ws.cell(r, 2).value
    vals = []
    for c in range(3, 27):
        v = ws.cell(r, c).value
        if v is None:
            vals.append(None)
        elif isinstance(v, (int, float)):
            vals.append(float(v))
        else:
            text_cells[(c, repr(v))] += 1
            vals.append(None)
    rows.append((r, a, b, vals))

# 结构：每4行一组
print('数据行数:', len(rows), '= 4 ×', len(rows) / 4)
bad = []
for i in range(0, len(rows), 4):
    blk = rows[i:i+4]
    if [x[2] for x in blk] != times:
        bad.append(('时刻序列', blk[0][0]))
    if blk[0][1] is None:
        bad.append(('首行缺日期', blk[0][0]))
    for x in blk[1:]:
        if x[1] is not None:
            bad.append(('非首行有日期', x[0]))
print('结构异常:', bad if bad else '无')

# 日期覆盖
def pdate(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    return datetime.datetime.strptime(str(v).strip(), '%Y-%m-%d').date()

d0 = pdate(rows[0][1])
days = [pdate(rows[i][1]) for i in range(0, len(rows), 4)]
print('起:', days[0], '止:', days[-1], '天数:', len(days))
print('日期连续:', all((days[i+1] - days[i]).days == 1 for i in range(len(days)-1)))
print('首日 = 2025-1-1:', days[0] == datetime.date(2025, 1, 1), ' 末日 = 2025-12-31:', days[-1] == datetime.date(2025, 12, 31))

# 数值扫描
allv = [v for _, _, _, vals in rows for v in vals]
nonnull = [v for v in allv if v is not None]
print('单元格总数 %d，空 %d，数值 %d，非数值文本 %d' % (len(allv), len(allv) - len(nonnull), len(nonnull), sum(text_cells.values())))
if text_cells:
    print('非数值文本样例:', list(text_cells.items())[:10])
print('取值范围: min %.2f  max %.2f' % (min(nonnull), max(nonnull)))
neg = [v for v in nonnull if v < 0]
print('负值个数:', len(neg))

# 夜间零值占比
for k in [1, 7, 13, 24]:
    col = [vals[k-1] for _, _, _, vals in rows if vals[k-1] is not None]
    print('  预报%2d小时: 均值 %9.1f  零值占比 %5.1f%%  最大 %8.1f' % (k, sum(col)/len(col), 100*sum(1 for v in col if v == 0)/len(col), max(col)))
