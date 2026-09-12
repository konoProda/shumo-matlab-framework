# -*- coding: utf-8 -*-
"""逐 (发布时刻,k) 的 MAE 诊断 + 全年平均日曲线对照"""
import openpyxl

wb2 = openpyxl.load_workbook('data/附件/附件2.xlsx', data_only=True)
ws2 = wb2['光伏发电实际功率']
D, T = 365, 144
raw = [[float(ws2.cell(r, c).value) for c in range(2, 2+T)] for r in range(2, 2+D)]
pv = [[raw[0][T-1]] + raw[0][:T-1]] + [[raw[d-1][T-1]] + raw[d-1][:T-1] for d in range(1, D)]
act = [[sum(pv[d][6*h:6*h+6])/6.0 for h in range(24)] for d in range(D)]

wb3 = openpyxl.load_workbook('data/附件/附件3.xlsx', data_only=True)
ws3 = wb3['Sheet1']
tmap = {'0:00': 0, '6:00': 6, '12:00': 12, '18:00': 18}
iss = [(ws3.cell(r, 2).value, [float(ws3.cell(r, c).value or 0) for c in range(3, 27)])
       for r in range(2, ws3.max_row + 1)]

def get(d, lab, k):
    return iss[d*4 + ['0:00', '6:00', '12:00', '18:00'].index(lab)][1][k-1]

print('=== 逐 (发布时刻, k) 的 MAE：读法A 用 act[τ+k-1]，读法B 用 act[τ+k] ===')
print('%-6s %-4s %10s %10s %10s' % ('发布', 'k', '读法A', '读法B', '谁更优'))
for lab in ['0:00', '6:00', '12:00', '18:00']:
    for k in (1, 2, 3, 6, 9, 12, 13, 18, 24):
        ea = eb = 0.0; n = 0
        for d in range(D):
            ha = tmap[lab] + k - 1
            if d + ha//24 >= D: continue
            a = act[d + ha//24][ha % 24]
            hb = tmap[lab] + k
            b = act[d + hb//24][hb % 24] if d + hb//24 < D else None
            if b is None: continue
            v = get(d, lab, k)
            ea += abs(v - a); eb += abs(v - b); n += 1
        print('%-6s %-4d %10.1f %10.1f %10s' % (lab, k, ea/n, eb/n, 'A' if ea < eb else 'B'))

print()
print('=== 全年平均日曲线（0:00 发布）===')
print('%4s %10s %12s %12s' % ('小时', '实际', '预报(读A)', '预报(读B)'))
for h in range(24):
    a = sum(act[d][h] for d in range(D))/D
    fa = sum(get(d, '0:00', h+1) for d in range(D-1))/ (D-1)          # k=h+1 ↔ 读A 的第 h 小时
    fb = sum(get(d, '0:00', h) for d in range(1, D))/ (D-1)            # k=h   ↔ 读B 的第 h 小时
    print('%4d %10.1f %12.1f %12.1f' % (h, a, fa, fb))
