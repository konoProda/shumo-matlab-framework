# -*- coding: utf-8 -*-
"""附件3 预报精度画像：按发布时刻/预报步长的 MAE，与问题二自建预测器对照"""
import openpyxl, json

wb2 = openpyxl.load_workbook('data/附件/附件2.xlsx', data_only=True)
ws2 = wb2['光伏发电实际功率']
D, T = 365, 144
raw = [[float(ws2.cell(r, c).value) for c in range(2, 2+T)] for r in range(2, 2+D)]
pv = [[raw[0][T-1]] + raw[0][:T-1]] + [[raw[d-1][T-1]] + raw[d-1][:T-1] for d in range(1, D)]
act = [[sum(pv[d][6*h:6*h+6])/6.0 for h in range(24)] for d in range(D)]

wb3 = openpyxl.load_workbook('data/附件/附件3.xlsx', data_only=True)
ws3 = wb3['Sheet1']
lab2i = {'0:00': 0, '6:00': 1, '12:00': 2, '18:00': 3}
tmap  = {'0:00': 0, '6:00': 6, '12:00': 12, '18:00': 18}
fy = [[[float(ws3.cell(d*4 + j + 2, c).value or 0) for c in range(3, 27)] for j in range(4)] for d in range(D)]

def err(issue, k):
    """返回 lead=k 的 MAE / RMSE（把所有 (d,τ) 汇总）"""
    e = []
    for d in range(D):
        for lab in ('0:00', '6:00', '12:00', '18:00'):
            hh = tmap[lab] + k - 1
            if hh >= 24: continue
            e.append(fy[d][lab2i[lab]][k-1] - act[d][hh])
    return sum(abs(x) for x in e)/len(e), (sum(x*x for x in e)/len(e))**0.5, sum(e)/len(e), len(e)

print('【附件3 精度画像】按预报步长 k（自发布时刻起第 k 个整点区间）')
print('%4s %10s %10s %10s %8s' % ('k', 'MAE(kW)', 'RMSE(kW)', '平均偏差', 'n'))
for k in (1,2,3,4,5,6,9,12,18,24):
    m, r, b, n = err(None, k)
    print('%4d %10.1f %10.1f %+10.1f %8d' % (k, m, r, b, n))

print()
print('【按发布时刻】k=1..24 汇总')
for lab in ('0:00', '6:00', '12:00', '18:00'):
    e = []
    for d in range(D):
        for k in range(1, 25):
            hh = tmap[lab] + k - 1
            if hh >= 24: continue
            e.append(fy[d][lab2i[lab]][k-1] - act[d][hh])
    print('  发布 %-5s  MAE %7.1f  RMSE %7.1f  偏差 %+7.1f  n=%d' % (
        lab, sum(abs(x) for x in e)/len(e), (sum(x*x for x in e)/len(e))**0.5, sum(e)/len(e), len(e)))

# 问题二自建预测器：最近 K 个同星期日的均值（严格早于当天），K=4
print()
print('【问题二自建预测器 K=4（同星期回溯均值）逐小时 MAE，与附件3 的 0:00 发布对照】')
K = 4
mae_self = mae_f3 = 0.0; n = 0
for d in range(28, D):                    # 前 4 周无足够历史
    r = d % 7
    S = [d - 7*(j+1) for j in range(K) if d - 7*(j+1) >= 0 and (d - 7*(j+1)) % 7 == r]
    for h in range(24):
        if not S: continue
        p = sum(act[s][h] for s in S)/len(S)
        mae_self += abs(p - act[d][h]); mae_f3 += abs(fy[d][0][h] - act[d][h]); n += 1
print('  自建预测器 MAE = %.1f kW   附件3(0:00发布) MAE = %.1f kW   （n=%d 小时点）' % (mae_self/n, mae_f3/n, n))
